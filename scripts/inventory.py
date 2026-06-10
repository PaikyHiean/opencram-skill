"""
inventory.py — Phase 0 盘点。

扫描用户指定的课件文件夹，分类 .ppt/.pptx/.pdf，尽力探测文字层，产出
work/manifest.json。

重要：本阶段（Phase 0）早于 bootstrap（Phase 3），LibreOffice 可能尚未安装，
因此**不依赖 LibreOffice**：
    - .pptx → python-pptx 直接探文字层与页数
    - .pdf  → pypdf 直接探文字层与页数
    - .ppt  → 旧二进制格式，无法在此廉价探测，标记 text_layer="unknown"，
              留到 extract 阶段（LibreOffice 转换后）确认；扫描件届时会被标红。

python-pptx / pypdf 若未安装，则优雅降级为 text_layer="unknown"。

用法：
    python scripts/inventory.py <课件文件夹>            # 默认递归
    python scripts/inventory.py <课件文件夹> --no-recursive
    python scripts/inventory.py <课件文件夹> --json
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from lib import common as C  # noqa: E402

SUPPORTED_EXTS = {".ppt", ".pptx", ".pdf"}
# 文字层判定阈值：去空白后总字符数低于此值视为"疑似无文字层"
TEXT_LAYER_MIN_CHARS = 30


def _probe_pptx(path: Path) -> dict:
    try:
        from pptx import Presentation  # type: ignore
    except Exception as e:  # noqa: BLE001
        return {"text_layer": "unknown", "slides": None,
                "note": f"python-pptx 不可用: {e}"}
    try:
        prs = Presentation(str(path))
        n = len(prs.slides)
        chars = 0
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    chars += len("".join(shape.text_frame.text.split()))
        has = chars >= TEXT_LAYER_MIN_CHARS
        return {"text_layer": "yes" if has else "suspect_none",
                "slides": n, "text_chars": chars}
    except Exception as e:  # noqa: BLE001
        return {"text_layer": "unknown", "slides": None, "note": f"读取失败: {e}"}


def _probe_pdf(path: Path) -> dict:
    try:
        from pypdf import PdfReader  # type: ignore
    except Exception as e:  # noqa: BLE001
        return {"text_layer": "unknown", "pages": None,
                "note": f"pypdf 不可用: {e}"}
    try:
        reader = PdfReader(str(path))
        n = len(reader.pages)
        chars = 0
        for page in reader.pages[: min(n, 10)]:  # 抽前 10 页判断即可
            try:
                chars += len("".join((page.extract_text() or "").split()))
            except Exception:  # noqa: BLE001
                continue
        has = chars >= TEXT_LAYER_MIN_CHARS
        return {"text_layer": "yes" if has else "suspect_none",
                "pages": n, "text_chars": chars}
    except Exception as e:  # noqa: BLE001
        return {"text_layer": "unknown", "pages": None, "note": f"读取失败: {e}"}


def scan(folder: Path, recursive: bool) -> dict:
    if not folder.exists() or not folder.is_dir():
        raise SystemExit(f"文件夹不存在或不是目录：{folder}")

    it = folder.rglob("*") if recursive else folder.glob("*")
    sources = []
    for p in sorted(it):
        if not p.is_file():
            continue
        ext = p.suffix.lower()
        if ext not in SUPPORTED_EXTS:
            continue
        entry: dict = {
            "name": p.name,
            "rel_path": str(p.relative_to(folder)),
            "abs_path": str(p.resolve()),
            "ext": ext,
            "size_bytes": p.stat().st_size,
        }
        if ext == ".pptx":
            entry.update(_probe_pptx(p))
        elif ext == ".pdf":
            entry.update(_probe_pdf(p))
        else:  # .ppt
            entry.update({
                "text_layer": "unknown",
                "slides": None,
                "note": "旧版 .ppt：将在抽取阶段经 LibreOffice 转换后确认文字层",
            })
        sources.append(entry)
        if len(sources) > 500:
            C.eprint("⚠ 已超过 500 个匹配文件，停止扫描。请检查路径是否正确（避免指向根目录）。")
            break

    warnings = []
    for s in sources:
        if s.get("text_layer") == "suspect_none":
            warnings.append(
                f"{s['name']} 疑似无文字层（可能是扫描件/纯图片）——本期不做 OCR，"
                f"抽取结果会很少，请确认是否仍要纳入。"
            )

    manifest = {
        "folder": str(folder.resolve()),
        "recursive": recursive,
        "count": len(sources),
        "by_ext": {ext: sum(1 for s in sources if s["ext"] == ext)
                   for ext in sorted(SUPPORTED_EXTS)},
        "sources": sources,
        "warnings": warnings,
    }
    return manifest


def _human(m: dict) -> str:
    lines = [f"=== 盘点结果：{m['folder']} ===",
             f"共 {m['count']} 个受支持文件  "
             f"({', '.join(f'{k} {v}' for k, v in m['by_ext'].items() if v)})",
             ""]
    for s in m["sources"]:
        n = s.get("slides") or s.get("pages")
        npart = f"{n}页/张" if n else "页数待定"
        tl = {"yes": "有文字层", "suspect_none": "⚠ 疑似无文字层",
              "unknown": "文字层待定"}.get(s.get("text_layer"), "?")
        lines.append(f"  • {s['name']}  [{s['ext']}]  {npart}  {tl}")
    if m["warnings"]:
        lines.append("")
        lines.append("⚠ 警告：")
        for w in m["warnings"]:
            lines.append(f"  - {w}")
    return "\n".join(lines)


def main() -> int:
    ap = argparse.ArgumentParser(description="开卷考 skill Phase 0 盘点")
    ap.add_argument("folder", help="课件所在文件夹路径")
    ap.add_argument("--no-recursive", dest="recursive", action="store_false",
                    help="只扫描该文件夹本层，不递归子目录")
    ap.add_argument("--json", action="store_true", help="仅输出 JSON")
    ap.set_defaults(recursive=True)
    args = ap.parse_args()

    C.ensure_dirs()
    manifest = scan(Path(args.folder), args.recursive)
    C.write_json(C.MANIFEST_PATH, manifest)

    if args.json:
        print(json.dumps(manifest, ensure_ascii=False, indent=2))
    else:
        print(_human(manifest))
        print(f"\n→ 已写入 {C.MANIFEST_PATH}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
